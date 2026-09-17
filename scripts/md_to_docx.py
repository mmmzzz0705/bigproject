# -*- coding: utf-8 -*-
"""把项目内的 Markdown 报告转成排版规范的 .docx。

用法：
    python scripts/md_to_docx.py <input.md> [output.docx]

特性：
- 标题 H1/H2/H3、正文、无序/有序列表
- 表格（含表头加粗底色、自动列宽）
- 围栏代码块（等宽字体 + 浅灰底纹，保留换行）
- 引用块（左侧竖线 + 浅灰底纹）
- 行内 **加粗** 与 `代码`
- 全文中文字体（正文宋体 / 标题微软雅黑）+ A4 页边距 + 页脚页码
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "宋体"
BODY_FONT_EN = "Times New Roman"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"
MONO_FONT_EA = "等线"

ACCENT = RGBColor(0x1F, 0x3B, 0x73)
GREY = RGBColor(0x59, 0x59, 0x59)
CODE_BG = "F2F4F7"
QUOTE_BG = "F7F8FA"
HEAD_BG = "DCE6F1"


# ---------------------------------------------------------------- 基础工具
def set_font(run, *, ascii_font=None, east=None, size=None, bold=None,
             italic=None, color=None):
    if ascii_font:
        run.font.name = ascii_font
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if ascii_font or east:
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        if ascii_font:
            rfonts.set(qn("w:ascii"), ascii_font)
            rfonts.set(qn("w:hAnsi"), ascii_font)
        if east:
            rfonts.set(qn("w:eastAsia"), east)
    return run


def shade(element, fill: str):
    """给段落或单元格加底纹。"""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element
    pr.append(shd)


def left_bar(paragraph, color="1F3B73", size=18):
    """段落左侧竖线（用于引用块）。"""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    ppr.append(borders)


def add_toc_field(doc):
    """插入可更新的目录域（Word 中按 F9 或右键“更新域”）。"""
    p = doc.add_paragraph()
    run = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    tip = OxmlElement("w:t")
    tip.text = "（在 Word 中右键选择“更新域”生成目录）"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(fld)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(tip)
    run._r.append(end)
    return p


def add_page_number_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run("— ")
    set_font(r1, ascii_font=BODY_FONT_EN, east=BODY_FONT, size=9, color=GREY)
    run = p.add_run()
    for el, attr in (("w:fldChar", "begin"), ("w:instrText", None),
                     ("w:fldChar", "separate"), ("w:t", None), ("w:fldChar", "end")):
        node = OxmlElement(el)
        if el == "w:fldChar":
            node.set(qn("w:fldCharType"), attr)
        elif el == "w:instrText":
            node.set(qn("xml:space"), "preserve")
            node.text = " PAGE "
        elif el == "w:t":
            node.text = "1"
        run._r.append(node)
    set_font(run, ascii_font=BODY_FONT_EN, east=BODY_FONT, size=9, color=GREY)
    r2 = p.add_run(" —")
    set_font(r2, ascii_font=BODY_FONT_EN, east=BODY_FONT, size=9, color=GREY)


# ---------------------------------------------------------------- 行内解析
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text, *, size=10.5, base_font=None, base_east=None,
               color=None, bold_all=False):
    """解析 **加粗** 与 `代码`，逐段写入 run。"""
    base_font = base_font or BODY_FONT_EN
    base_east = base_east or BODY_FONT
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            set_font(paragraph.add_run(part[2:-2]), ascii_font=base_font,
                     east=base_east, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            set_font(paragraph.add_run(part[1:-1]), ascii_font=MONO_FONT,
                     east=MONO_FONT_EA, size=size - 0.5,
                     color=RGBColor(0xC0, 0x39, 0x2B))
        else:
            set_font(paragraph.add_run(part), ascii_font=base_font,
                     east=base_east, size=size, bold=bold_all, color=color)


# ---------------------------------------------------------------- 文档构建
def setup_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.6)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = BODY_FONT_EN
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.5

    sizes = {1: 18, 2: 14, 3: 12}
    for lvl in (1, 2, 3):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = HEAD_FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        st.font.size = Pt(sizes[lvl])
        st.font.bold = True
        st.font.color.rgb = ACCENT if lvl < 3 else RGBColor(0x2E, 0x5A, 0x88)
        st.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.25
        st.paragraph_format.keep_with_next = True
    return doc


TABLE_RE = re.compile(r"^\s*\|(.+)\|\s*$")
SEP_RE = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")


def split_row(line):
    inner = line.strip().strip("|")
    return [c.strip() for c in inner.split("|")]


def build_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    hdr = table.rows[0]
    for i, text in enumerate(header):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        add_inline(p, text, size=9.5, bold_all=True, color=ACCENT)
        shade(cell._tc, HEAD_BG)

    for row in body:
        cells = table.add_row().cells
        for i in range(ncols):
            text = row[i] if i < len(row) else ""
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, text, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def convert(md_path: Path, docx_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = setup_doc()
    add_page_number_footer(doc)

    i, n = 0, len(lines)
    inserted_toc = False
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 空行
        if not stripped:
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "C9CFDA")
            bdr.append(bottom)
            pPr.append(bdr)
            i += 1
            continue

        # 围栏代码块
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            for code_line in buf:
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.space_after = Pt(0)
                pf.space_before = Pt(0)
                pf.line_spacing = 1.0
                pf.left_indent = Cm(0.4)
                set_font(p.add_run(code_line if code_line else " "),
                         ascii_font=MONO_FONT, east=MONO_FONT_EA, size=9)
                shade(p._p, CODE_BG)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 表格
        if TABLE_RE.match(line) and i + 1 < n and SEP_RE.match(lines[i + 1] or ""):
            rows = []
            j = i
            while j < n and TABLE_RE.match(lines[j] or ""):
                if not SEP_RE.match(lines[j]):
                    rows.append(split_row(lines[j]))
                j += 1
            if rows:
                build_table(doc, rows)
            i = j
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                h = doc.add_heading("", level=0)
                pf = h.paragraph_format
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf.space_before = Pt(0)
                pf.space_after = Pt(6)
                set_font(h.add_run(text), ascii_font=HEAD_FONT, east=HEAD_FONT,
                         size=22, bold=True, color=ACCENT)
                h.style.paragraph_format.keep_with_next = True
            else:
                h = doc.add_heading("", level=min(level, 4))
                add_inline(h, text, size={2: 15, 3: 12.5, 4: 11}[level],
                           base_font=HEAD_FONT, base_east=HEAD_FONT,
                           bold_all=True,
                           color=ACCENT if level == 2 else RGBColor(0x2E, 0x5A, 0x88))
            # 首个一级标题之后插入目录
            if level == 1 and not inserted_toc:
                sub = doc.add_paragraph()
                sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(sub.add_run("（目录见下页，在 Word 中右键“更新域”即可生成）"),
                         ascii_font=BODY_FONT, east=BODY_FONT, size=9,
                         color=GREY, italic=True)
                doc.add_page_break()
                toc_h = doc.add_heading("", level=1)
                add_inline(toc_h, "目录", size=18, base_font=HEAD_FONT,
                           base_east=HEAD_FONT, bold_all=True, color=ACCENT)
                add_toc_field(doc)
                doc.add_page_break()
                inserted_toc = True
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            for ql in buf:
                if not ql:
                    continue
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Cm(0.6)
                pf.space_before = Pt(2)
                pf.space_after = Pt(2)
                pf.line_spacing = 1.35
                add_inline(p, ql, size=9.5, color=GREY)
                shade(p._p, QUOTE_BG)
                left_bar(p)
            continue

        # 列表
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            pf = p.paragraph_format
            pf.left_indent = Cm(0.75)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.4
            add_inline(p, m.group(2).strip(), size=10.5)
            i += 1
            continue

        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.75)
            pf.first_line_indent = Cm(-0.55)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.4
            add_inline(p, f"{m.group(2)}. {m.group(3).strip()}", size=10.5)
            i += 1
            continue

        # 普通段落（支持硬换行续行）
        buf = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", ">", "-", "*", "```", "|"))
                    or re.match(r"^\d+[.)]\s", nxt)
                    or re.fullmatch(r"-{3,}", nxt)):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        add_inline(p, "".join(buf), size=10.5)
        continue

    doc.save(str(docx_path))
    return docx_path


def main():
    if len(sys.argv) < 2:
        print("用法: python md_to_docx.py <input.md> [output.docx]")
        return 2
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    out = convert(src, dst)
    print(f"OK -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
