# -*- coding: utf-8 -*-
"""按《需求规格说明书（SRS）.docx》模板规模，生成《政明白》SRS。

用法：
  <venv python> D:\\group6\\scripts\\build_srs.py

要点：
1. 只做"占位符 → 政明白真实内容"的替换，不新增章节、不新增表格；
2. 3 个功能点（智能问答 / 会话与历史 / 语料管理）由模板 3.1 块克隆而来；
3. 模板自带的第二个子功能点块（FR3.1.3.2 三行）必须删除，否则残留占位符。
"""
from copy import deepcopy

from docx import Document
from docx.text.paragraph import Paragraph

SRC = r"D:\cxdownload\需求规格说明书（SRS）.docx"
OUT = r"D:\cxdownload\政明白_需求规格说明书（SRS）_V1.0.docx"

doc = Document(SRC)
t_terms, t_user, t_tbd, t_hist, t_appr = doc.tables[:5]


def find(text, startswith=False):
    for p in doc.paragraphs:
        t = p.text.strip()
        if (t.startswith(text) if startswith else text in t):
            return p
    raise KeyError("未找到段落：" + text)


def set_text(p, text):
    """替换段落文本，保留首个 run 的字体；\n 转成软换行。"""
    lines = text.split("\n")
    runs = p.runs
    if not runs:
        p.add_run(lines[0])
        runs = p.runs
    else:
        runs[0].text = lines[0]
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    tmpl = runs[0]
    for line in lines[1:]:
        rb = p.add_run()
        if tmpl._element.rPr is not None:
            rb._element.insert(0, deepcopy(tmpl._element.rPr))
        rb.add_break()
        rn = p.add_run(line)
        if tmpl._element.rPr is not None:
            rn._element.insert(0, deepcopy(tmpl._element.rPr))


def set_cell(cell, text, bold=False):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    set_text(cell.paragraphs[0], text)
    if bold:
        for r in cell.paragraphs[0].runs:
            r.bold = True


def add_row_like(table, src_idx=1):
    table._tbl.append(deepcopy(table.rows[src_idx]._tr))
    return table.rows[-1]


def clone_paras(src_paras, after):
    """把 src_paras 依次深拷贝插到 after 之后，返回新段落列表。"""
    cur = after
    out = []
    for sp in src_paras:
        el = deepcopy(sp._element)
        cur._element.addnext(el)
        cur = Paragraph(el, sp._parent)
        out.append(cur)
    return out


# ---------------- 封面 ----------------
set_text(find("项目名称："), "\n".join([
    "项目名称： 政明白（GovRAG 政务智能问答与办事引导系统）",
    "项目小组： 第六组",
    "文档编号： GOVRAG-SRS-V1.0",
    "项目经理： 麦少彤",
    "需求分析人员： 陈梓烽、贺悦洋、黎敬浩",
    "创建日期： 2026-09-18",
    "最后更新日期： 2026-09-18",
    "文档状态： 评审中",
]))

# ---------------- 1 引言 ----------------
set_text(find("本文档旨在明确"), "本文档旨在明确《政明白》（GovRAG 政务智能问答与办事引导系统）的功能和非功能性需求，为开发团队、测试团队及项目相关方提供共同的理解基础，并作为项目设计、开发、测试和验收的依据。")
set_text(find("项目目标："), "项目目标： 构建一个严格依据已入库政务办事指南作答的检索增强问答（RAG）系统——对已收录事项输出材料清单、办理流程、受理部门、办结时限与溯源信息，对未收录事项明确拒答并引导至线下窗口或 12345 热线。")
set_text(find("主要功能："), "主要功能： 政务事项问答；结构化办事材料清单；答案溯源；相关性闸门与拒答；多轮会话与历史管理；语料自助增删（工作台）；链路健康自检。")
set_text(find("不包含内容："), "不包含内容： 不提供在线申报与审批办理，不做用户身份认证与电子证照核验，不接入统一政务事项数据源；事项范围仅限已入库的 33 篇广东政务办事指南。")

term_rows = [
    ("RAG", "检索增强生成：生成答案前先从知识库检索相关片段作为上下文，使输出可溯源"),
    ("混合检索", "向量检索（语义）与 BM25 检索（关键词）并行后按 0.45×向量分 + 0.55×关键词分融合"),
    ("相关性闸门", "向量分 ≥0.45，或关键词分 ≥44.3 且融合分 ≥0.62 方予放行，未通过则直接拒答"),
]
set_cell(t_terms.rows[4].cells[0], term_rows[0][0])
set_cell(t_terms.rows[4].cells[1], term_rows[0][1])
for k, v in term_rows[1:]:
    row = add_row_like(t_terms, 4)
    set_cell(row.cells[0], k)
    set_cell(row.cells[1], v)

for p, t in zip(
    [find("[可行性分析报告文档编号及名称]"), find("[用户需求调研记录或相关会议纪要]"), find("[相关的行业标准或规范]")],
    [
        "《第六组_政明白_可行性分析报告》V1.1，2026 年 9 月 13 日",
        "《软件工程与项目案例教程》，梁立新、郭锐，清华大学出版社",
        "IEEE Std 830-1998, IEEE Recommended Practice for Software Requirements Specifications [S]；广东政务服务网办事指南（https://www.gdzwfw.gov.cn/）",
    ],
):
    set_text(p, t)

set_text(find("本文档主要包括总体描述"), "本文档主要包括总体描述、功能性需求、外部接口需求、非功能性需求和其他需求等章节。")

# ---------------- 2 总体描述 ----------------
set_text(find("[描述软件产品的背景"), "政明白是面向办事群众的政务智能问答系统。当前群众办事前需咨询“要什么材料、去哪办、几天办好”，只能依赖窗口人工答复或自行翻阅办事指南（单篇 1,928~45,082 字），存在查阅成本高、同类事项易混淆、答复口径不统一三类问题。本系统作为窗口人工咨询的辅助手段独立建设，内容来源为广东政务服务网公开办事指南并以引用方式标注出处，无依据时引导至线下窗口或 12345 热线，不替换任何现有办理系统。")

for i, (a, b, c) in enumerate([
    ("办事群众", "需办事前咨询材料、流程、时限的普通市民，无系统使用经验，习惯口语化提问", "自然语言提问即可获得准确、可核验的答复"),
    ("政务窗口工作人员", "熟悉业务、需快速核对答复口径的一线人员", "快速查证材料清单与办结时限，口径与官方指南一致"),
    ("知识库管理员（小组成员兼任）", "负责语料入库、维护与链路检查，不参与代码修改", "上传/粘贴/查看/删除语料，查看链路状态"),
]):
    set_cell(t_user.rows[i + 1].cells[0], a)
    set_cell(t_user.rows[i + 1].cells[1], b)
    set_cell(t_user.rows[i + 1].cells[2], c)

set_text(find("硬件环境："), "硬件环境： 服务端运行于 Docker Desktop 容器（建议 8 GB 以上可用内存、20 GB 以上可用磁盘）；客户端为普通 PC 或移动端浏览器，无特殊硬件要求。")
set_text(find("软件环境："), "软件环境： 后端 Python 3.12 + FastAPI + Uvicorn（8000）；前端 Vue 3 + Vite 构建、Nginx 托管（8080）；数据库 PostgreSQL 16（5432）；向量库 Milvus 2.x（19531）；大模型与 Embedding 调用阿里云百炼 DashScope。客户端为 Chrome / Edge / Firefox 最新版本。")
set_text(find("网络环境："), "网络环境： 容器间经 Docker 自定义网络通信；浏览器经 Nginx 反向代理访问 /api；访问 DashScope 需公网 HTTPS，建议带宽 ≥10 Mbps。")

set_text(find("[必须使用的特定技术或平台]"), "技术平台： Python 3.12 + FastAPI（后端）、Vue 3 + Vite（前端）、Milvus 向量库（1024 维）、PostgreSQL 16，Docker Compose 一键部署。")
set_text(find("[必须遵循的行业标准或规范]"), "标准规范： 遵循 IEEE Std 830-1998 与 RESTful 接口约定；政务内容以广东政务服务网官方发布文本为准。")
set_text(find("[必须兼容的已有系统或接口]"), "兼容要求： 语料源自广东政务服务网办事指南 PDF；Embedding 与生成必须调用 DashScope 的 /compatible-mode/v1 兼容接口；前端请求经 Nginx 转发至后端。")
set_text(find("[法律法规和政策限制]"), "法规与政策： 仅引用公开政务信息并标注出处；不采集、不外传用户隐私数据；不做身份认证与电子证照核验。")

set_text(find("[假设用户具备一定的计算机操作能力]"), "假设用户具备基本的中文读写与浏览器操作能力，无需培训即可提问；单条问题不超过 1000 字。")
set_text(find("[假设网络连接基本稳定]"), "假设网络连接基本稳定且可访问 DashScope；办事指南内容准确、具备时效性，失效时需重新导入。")
set_text(find("[项目成功依赖的第三方服务或组件]"), "项目依赖 DashScope 大模型与 Embedding 服务可持续访问、额度充足；依赖 Milvus 与 PostgreSQL 正常启动；依赖相关事项的办事指南已完成入库。")

# ---------------- 3 功能性需求 ----------------
p_31 = find("3.1 功能点")
p_311 = find("3.1.1 描述与优先级")
p_312 = find("3.1.2 激发/响应序列")
p_exc = find("功能点激发（用户动作）")
p_res = find("响应（系统反应）")
p_313 = find("3.1.3 子功能点列表")
p_fr1 = find("FR3.1.3.1：")
p_src1 = find("需求来源：", startswith=True)
p_acc1 = Paragraph(p_src1._element.getnext(), p_src1._parent)
FR_BLOCK = [p_fr1, p_src1, p_acc1]

# 删除模板自带的第二块（FR3.1.3.2 三行）+ 两行提示语，后续由克隆补齐
_p = find("FR3.1.3.2：")
_e, _els = _p._element, [_p._element]
for _ in range(2):
    _e = _e.getnext()
    _els.append(_e)
for _e in _els:
    _e.getparent().remove(_e)
for _p in (find("[更多功能需求...]"), find("[按此格式描述系统的每一个功能模块...]")):
    _p._element.getparent().remove(_p._element)


def fill_fp(title, desc, prio, exc, resp, frs, no):
    prefix = no.rsplit(".", 1)[0]
    set_text(p_31, title)
    set_text(p_311, prefix + ".1 描述与优先级： " + desc + "\n优先级： " + prio)
    set_text(p_312, prefix + ".2 激发/响应序列")
    set_text(p_exc, exc)
    set_text(p_res, resp)
    set_text(p_313, prefix + ".3 子功能点列表")
    blocks = [FR_BLOCK]
    for _ in range(len(frs) - 1):
        blocks.append(clone_paras(FR_BLOCK, blocks[-1][-1]))
    for i, blk in enumerate(blocks):
        set_text(blk[0], "FR%s.%d： %s" % (no, i + 1, frs[i][0]))
        set_text(blk[1], "需求来源： %s" % frs[i][1])
        set_text(blk[2], "验收标准： %s" % frs[i][2])
    return blocks[-1][-1]


def clone_fp(after, title, desc, prio, exc, resp, frs, no):
    src = [p_31, p_311, p_312, p_exc, p_res, p_313, FR_BLOCK[0], FR_BLOCK[1], FR_BLOCK[2]]
    new = clone_paras(src, after)
    prefix = no.rsplit(".", 1)[0]
    set_text(new[0], title)
    set_text(new[1], prefix + ".1 描述与优先级： " + desc + "\n优先级： " + prio)
    set_text(new[2], prefix + ".2 激发/响应序列")
    set_text(new[3], exc)
    set_text(new[4], resp)
    set_text(new[5], prefix + ".3 子功能点列表")
    blocks = [[new[6], new[7], new[8]]]
    for _ in range(len(frs) - 1):
        blocks.append(clone_paras(blocks[-1], blocks[-1][-1]))
    for i, blk in enumerate(blocks):
        set_text(blk[0], "FR%s.%d： %s" % (no, i + 1, frs[i][0]))
        set_text(blk[1], "需求来源： %s" % frs[i][1])
        set_text(blk[2], "验收标准： %s" % frs[i][2])
    return blocks[-1][-1]


last = fill_fp(
    "3.1 智能问答（RAG 核心链路）",
    "接收用户的政务办事咨询，经问题清洗、向量化、混合检索、相关性闸门判定与大模型生成，输出可溯源的结构化答复；未通过闸门时直接拒答并引导，不调用大模型。",
    "高",
    "功能点激发（用户动作）： 用户在输入框输入办事咨询问题，点击发送或按 Enter。",
    "响应（系统反应）： 系统在 30 秒内返回答案、材料清单与溯源列表；未通过闸门时在 0.5 秒内返回拒答话术与引导。",
    [
        ("混合检索与相关性闸门：Milvus 向量检索与 BM25 关键词检索并行取 Top-K=4；向量分 ≥0.45，或关键词分 ≥44.3 且融合分 ≥0.62 方予放行",
         "业务规则（杜绝幻觉）",
         "域外 5 例与领域内未收录 4 例全部被拦截且不调用大模型；域内 18 例全部放行"),
        ("结构化输出与溯源：返回答案、材料清单（必备/可选/流程/提示/部门/时限/收费）与来源列表（文档名、章节、片段、相似度）",
         "用户需求（答案可核验）",
         "每条答案 sources 非空，相似度保留 3 位小数，点击可查看命中片段"),
    ], "3.1.3")

last = clone_fp(
    last, "3.2 会话与历史管理",
    "为每位用户提供独立会话，支持新建、切换、清空与删除，并在多轮问答中按窗口裁剪历史上下文。",
    "中",
    "功能点激发（用户动作）： 用户新建、切换、清空或删除会话。",
    "响应（系统反应）： 后端返回 session_id 或该会话的历史问答列表，前端同步刷新会话列表与主区消息。",
    [
        ("创建会话与自动补建：POST /api/session/create 返回 session_id；会话在库中不存在时自动补建，不返回 500",
         "用户需求（会话隔离）",
         "清库后使用旧 session_id 提问仍能正常作答"),
        ("历史裁剪与查询：保留最近 5 轮、正文裁剪至 3000 字符以内；GET /api/chat/history 按时间升序返回问题、答案与时间",
         "技术需求（成本与延迟控制）",
         "切换会话后完整还原该会话历史，顺序与时间正确"),
    ], "3.2.3")

clone_fp(
    last, "3.3 语料管理（工作台）",
    "面向知识库管理员提供语料的自助入库与维护：列表统计、文件上传、文本粘贴、切分预览、删除与重新导入。",
    "高",
    "功能点激发（用户动作）： 管理员上传文件、粘贴文本、查看切分预览、删除语料或点击重新导入。",
    "响应（系统反应）： 系统完成入库或删除，返回受影响片段数与文档信息，并刷新统计与语料列表。",
    [
        ("上传与粘贴入库：文件 ≤10 MB 落盘至 data/uploads（与内置语料同名自动改名），粘贴正文 20~200,000 字；按 500 字切分、批量向量化后幂等写入",
         "用户需求（自助扩容）",
         "同一文件上传两次片段总数不变；与内置语料同名时自动改名，基准语料不被覆盖"),
        ("删除与重新导入：上传语料连文件与元数据一并删除；内置语料只清向量并置为未入库，可随时 reingest",
         "业务规则（可恢复性）",
         "删除内置语料后列表中该篇 in_index=false 且可重新导入，片段数与原入库一致"),
    ], "3.3.3")

# ---------------- 4 外部接口需求 ----------------
set_text(find("风格要求："), "风格要求： 简洁现代的政务风格，浅色卡片布局；气泡式对话，答案支持 Markdown，材料清单以卡片展示，侧栏提供会话列表与视图切换。")
set_text(find("适配要求："), "适配要求： 支持 Chrome、Edge、Firefox 最新两个版本，分辨率不低于 1280×720；前端构建产物由 Nginx 托管。")
set_text(find("UI原型："), "UI原型： 无独立原型文档，以已实现组件为准（Composer、MessageBubble、MaterialList、SourceRefs、Workbench 等）。")
set_text(find("[描述软件与硬件设备的所有交互"), "无特殊硬件接口需求：不依赖摄像头、读卡器或打印设备，全部交互通过浏览器完成。")
set_text(find("数据库："), "数据库： PostgreSQL 16（容器 gov-postgres:5432，库名 gov_qa），经 SQLAlchemy ORM 访问；缺失时降级为 SQLite。")
set_text(find("第三方API："), "第三方API： 阿里云百炼 DashScope OpenAI 兼容接口——Embedding 用 qwen3.7-text-embedding-flash（1024 维），生成用 deepseek-v4.1-flash（温度 0.2）。")
set_text(find("操作系统接口："), "操作系统接口： 无特殊要求，服务经 Docker 容器以 HTTP 方式对外提供。")
set_text(find("协议："), "协议： HTTP/1.1；调用 DashScope 走 HTTPS。")
set_text(find("数据格式："), "数据格式： JSON；语料上传为 multipart/form-data。")
set_text(find("API风格："), "API风格： RESTful，自带 /docs 交互式文档；主要接口为 POST /api/chat、POST /api/session/create、GET /api/chat/history、/api/corpus 系列、GET /api/health。")

# ---------------- 5 非功能性需求 ----------------
set_text(find("响应时间："), "响应时间： 标准问法 17~22 秒，口语化问法约 11 秒；被相关性闸门拦截的域外与未收录问题 0.3~0.4 秒。")
set_text(find("并发用户数："), "并发用户数： 支持至少 20 个用户同时在线提问；单轮问答仅产生 1 次 Embedding 与 1 次生成调用。")
set_text(find("数据量："), "数据量： 33 篇广东政务办事指南约 853 个文本片段，向量维度 1024；单条问题不超过 1000 字。")
set_text(find("用户密码需加密存储。"), "API 密钥仅保存在 .env 中，不入库、不写入文档、不返回前端；异常响应不暴露堆栈与内部路径。")
set_text(find("对不同角色的用户实行基于权限的访问控制。"), "访问控制： 工作台写操作支持令牌校验（后端 CORPUS_WRITE_TOKEN 与请求头 X-Workbench-Token 配对），配置后无令牌或错误令牌一律 403。")
set_text(find("接口需防范常见的Web攻击（如SQL注入、XSS）。"), "注入与脚本防护： 输入去除控制字符并限制长度，数据库访问全部走 ORM 参数化查询，前端不执行用户提供的脚本。")
set_text(find("系统可用性不低于99.9%。"), "降级可用性： 外部依赖（PostgreSQL、Milvus、大模型）任一缺失时系统仍可启动并提供服务，仅功能降级，不出现白屏或崩溃。")
set_text(find("系统平均无故障运行时间（MTBF）大于720小时。"), "就绪与恢复： 后端启动后 15~30 秒完成就绪（容器健康检查约 20 秒转 healthy）；依赖恢复后重启即可回到完整链路。")
set_text(find("界面直观，新用户无需培训即可完成基本操作。"), "界面直观： 新用户无需培训即可完成提问；首屏提供示例问题一键提问；答案附溯源信息可供自行核验。")
set_text(find("提供完整的在线帮助文档或操作指引。"), "即时反馈： 提供“重新生成”“点赞/点踩”等反馈入口；被拒答时说明原因并给出下一步建议。")
set_text(find("代码有良好的注释和文档。"), "代码分层： 后端按 routers / services / utils 分层，配置集中在 config.py 且全部可由环境变量覆盖。")
set_text(find("系统模块化设计，便于后期功能扩展和bug修复。"), "质量门禁： 提交前运行 scripts/quality_gate.py（ruff + pytest），配置一致性由 tests/test_env_sync.py 校验。")
set_text(find("[系统应能部署在主流Linux发行版上]"), "全部组件容器化，docker-compose 一键启停，可在 Windows（Docker Desktop）与主流 Linux 发行版部署。")
set_text(find("兼容性："), "兼容性： 兼容 Chrome、Edge、Firefox 最新两个版本；语料支持增量导入，新增事项需重新标定闸门阈值。")

# ---------------- 6 其他需求与附录 ----------------
set_text(find("《用户操作手册》"), "《用户操作手册》：说明提问、查看材料清单与溯源的操作方式。")
set_text(find("《系统安装部署手册》"), "《系统安装部署手册》：说明 docker-compose 启停、环境变量配置、语料重灌与常见故障排查。")
set_text(find("《API接口文档》"), "《API接口文档》：由 FastAPI 自动生成（/docs），含全部接口的请求与响应模型。")

for i, row_data in enumerate([
    ("语料仅覆盖 33 个广东政务事项，未收录事项是否接入统一政务数据源", "决定覆盖范围与拒答率", "麦少彤", "后续迭代确定"),
    ("公网部署时是否启用工作台访问令牌", "决定语料被误删或恶意清空的风险", "黎敬浩", "部署前确定"),
]):
    for j, v in enumerate(row_data):
        set_cell(t_tbd.rows[i + 1].cells[j], v)

for i, v in enumerate(["V1.0", "2026-09-18", "陈梓烽、贺悦洋、黎敬浩", "初稿创建"]):
    set_cell(t_hist.rows[1].cells[i], v)
for i, v in enumerate(["V1.1", "待评审后填写", "陈梓烽", "依据评审意见修订"]):
    set_cell(t_hist.rows[2].cells[i], v)

for i, (name, date) in enumerate([("陈梓烽、贺悦洋、黎敬浩", "2026-09-18"), ("陈梓烽", "2026-09-18"), ("贺悦洋", "2026-09-18"), ("麦少彤", "2026-09-18")]):
    set_cell(t_appr.rows[i + 1].cells[1], name)
    set_cell(t_appr.rows[i + 1].cells[3], date)

doc.save(OUT)

# ---------------- 自检 ----------------
bad = [p.text.strip() for p in Document(OUT).paragraphs if "[" in p.text and "IEEE" not in p.text]
print("saved:", OUT)
print("残留占位符：", bad if bad else "无")
